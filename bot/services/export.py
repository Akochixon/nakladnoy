from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import joinedload, selectinload

from bot.database.models import (
    Customer, Expediter, Invoice, InvoiceDuplicate, InvoiceItem, Product, SalesAgent, Supplier, TelegramUser,
)
from bot.services.stats import filters
from bot.utils import local


def _f(x):
    return float(x) if x is not None else None


def _sheet(wb: Workbook, title: str, headers: list[str], rows, widths: list[int], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    for r in rows:
        ws.append(r)
    return ws


async def _load_invoices(session: AsyncSession, since=None, user_pk=None) -> list[Invoice]:
    q = (
        select(Invoice)
        .options(
            joinedload(Invoice.customer), joinedload(Invoice.supplier),
            joinedload(Invoice.sales_agent), joinedload(Invoice.expediter), joinedload(Invoice.submitter),
            selectinload(Invoice.items).joinedload(InvoiceItem.product),
        )
        .where(*filters(since, user_pk))
        .order_by(Invoice.date.desc(), Invoice.id.desc())
    )
    return (await session.execute(q)).scalars().all()


def _invoice_sheets(wb: Workbook, invoices: list[Invoice]) -> None:
    inv_rows = [[
        i.invoice_number, i.date, i.supplier.name, i.customer.name,
        i.sales_agent.full_name if i.sales_agent else "", i.expediter.full_name if i.expediter else "",
        i.submitter.full_name, float(i.total_qty), float(i.total_sum),
        local(i.created_at).strftime("%d.%m.%Y %H:%M"), "bor" if i.photo_file_id else "yo'q",
    ] for i in invoices]
    item_rows = [[i.invoice_number, i.date, i.customer.name, it.product.code, it.product.name,
                  float(it.quantity), float(it.unit_price), float(it.line_total)]
                 for i in invoices for it in i.items]
    _sheet(wb, "Nakladnoylar", ["№", "Sana", "Postavshik", "Xaridor", "Agent", "Ekspeditor", "Yubordi",
                                "Jami soni", "Jami summa", "Kiritilgan vaqt", "Rasm"],
           inv_rows, [16, 12, 24, 30, 24, 22, 24, 12, 16, 18, 8], first=True)
    _sheet(wb, "Tovar qatorlari", ["Nakladnoy №", "Sana", "Xaridor", "Kod", "Tovar", "Soni", "Narx", "Summa"],
           item_rows, [16, 12, 30, 10, 44, 10, 14, 16])
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
            row[0].number_format = "DD.MM.YYYY"


async def build_invoices_xlsx(session: AsyncSession, since=None, user_pk=None) -> tuple[bytes, int]:
    invoices = await _load_invoices(session, since, user_pk)
    wb = Workbook()
    _invoice_sheets(wb, invoices)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(invoices)


async def _dim(session: AsyncSession, model, fk):
    q = (select(model, func.count(Invoice.id), func.coalesce(func.sum(Invoice.total_sum), 0))
         .outerjoin(Invoice, fk == model.id).group_by(model.id).order_by(model.id))
    return (await session.execute(q)).all()


async def build_full_xlsx(session: AsyncSession) -> tuple[bytes, dict[str, int]]:
    """Butun bazani Excelga chiqaradi: har bir jadval alohida varaq."""
    wb = Workbook()
    invoices = await _load_invoices(session)
    _invoice_sheets(wb, invoices)
    counts = {"Nakladnoylar": len(invoices)}

    async def add(title, headers, rows, widths):
        _sheet(wb, title, headers, rows, widths)
        counts[title] = len(rows)

    rows = await _dim(session, Customer, Invoice.customer_id)
    await add("Xaridorlar", ["ID", "Nomi", "INN", "Manzil", "Mijoz kodi", "Nakladnoylar", "Jami summa"],
              [[c.id, c.name, c.inn, c.address, c.client_code, n, float(t)] for c, n, t in rows],
              [6, 34, 14, 40, 16, 14, 16])
    rows = await _dim(session, Supplier, Invoice.supplier_id)
    await add("Postavshiklar", ["ID", "Nomi", "INN", "Manzil", "Telefon", "Nakladnoylar", "Jami summa"],
              [[s.id, s.name, s.inn, s.address, s.phone, n, float(t)] for s, n, t in rows],
              [6, 34, 14, 40, 16, 14, 16])
    rows = await _dim(session, SalesAgent, Invoice.sales_agent_id)
    await add("Agentlar", ["ID", "Ism-familya", "Nakladnoylar", "Jami summa"],
              [[a.id, a.full_name, n, float(t)] for a, n, t in rows], [6, 34, 14, 16])
    rows = await _dim(session, Expediter, Invoice.expediter_id)
    await add("Ekspeditorlar", ["ID", "Ism-familya", "Nakladnoylar", "Jami summa"],
              [[e.id, e.full_name, n, float(t)] for e, n, t in rows], [6, 34, 14, 16])

    q = (select(Product, func.count(func.distinct(InvoiceItem.invoice_id)),
                func.coalesce(func.sum(InvoiceItem.quantity), 0), func.coalesce(func.sum(InvoiceItem.line_total), 0),
                func.min(InvoiceItem.unit_price), func.max(InvoiceItem.unit_price))
         .outerjoin(InvoiceItem, InvoiceItem.product_id == Product.id).group_by(Product.id).order_by(Product.code))
    rows = (await session.execute(q)).all()
    await add("Tovarlar", ["ID", "Kod", "Nomi", "Birlik", "Nakladnoylar", "Jami soni", "Jami summa", "Min narx", "Max narx"],
              [[p.id, p.code, p.name, p.unit, n, float(qty), float(t), _f(lo), _f(hi)] for p, n, qty, t, lo, hi in rows],
              [6, 10, 44, 8, 14, 12, 16, 14, 14])

    q = (select(TelegramUser, func.count(Invoice.id), func.coalesce(func.sum(Invoice.total_sum), 0))
         .outerjoin(Invoice, Invoice.submitted_by == TelegramUser.id).group_by(TelegramUser.id).order_by(TelegramUser.id))
    rows = (await session.execute(q)).all()
    await add("Foydalanuvchilar",
              ["ID", "Ism-familya", "JSHSHIR", "Telegram ID", "Username", "Holat", "Rol", "Ro'yxatdan o'tgan",
               "Nakladnoylar", "Jami summa"],
              [[u.id, u.full_name, u.id_number, u.telegram_id, u.username, u.status.value, u.role,
                local(u.created_at).strftime("%d.%m.%Y %H:%M"), n, float(t)] for u, n, t in rows],
              [6, 28, 16, 14, 18, 12, 8, 18, 14, 16])

    q = (select(Invoice.invoice_number, TelegramUser.full_name, InvoiceDuplicate.decision, InvoiceDuplicate.created_at)
         .select_from(InvoiceDuplicate).join(Invoice, Invoice.id == InvoiceDuplicate.invoice_id)
         .join(TelegramUser, TelegramUser.id == InvoiceDuplicate.submitted_by)
         .order_by(InvoiceDuplicate.created_at.desc()))
    rows = (await session.execute(q)).all()
    await add("Dublikat loglari", ["Nakladnoy №", "Kim topshirgan", "Qaror", "Vaqt"],
              [[n, u, d.value, local(t).strftime("%d.%m.%Y %H:%M")] for n, u, d, t in rows], [16, 28, 24, 18])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), counts
